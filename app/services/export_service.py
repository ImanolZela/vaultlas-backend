from io import BytesIO
from typing import Optional
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Table, TableStyle,
    Paragraph, Spacer,
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Palette (light theme) ─────────────────────────────────────────────────────
NEON        = colors.HexColor("#CCFF00")
NEON_DARK   = colors.HexColor("#8AAA00")   # neon toned down for light bg text
INK         = colors.HexColor("#111111")   # near-black for main text
INK_MID     = colors.HexColor("#555555")   # secondary text
INK_LIGHT   = colors.HexColor("#999999")   # tertiary / labels
RULE        = colors.HexColor("#E5E5E5")   # hairline dividers
BG_PAGE     = colors.white                 # document background
BG_HEADER   = colors.HexColor("#F9F9F9")   # header area fill
BG_STATS    = colors.HexColor("#F4F4F4")   # stats bar
BG_ROW_ALT  = colors.HexColor("#FAFAFA")   # alternating row
BG_ROW_MAIN = colors.white
BG_FOOT     = colors.HexColor("#F4F4F4")
EMERALD     = colors.HexColor("#059669")   # slightly deeper for light bg
CORAL       = colors.HexColor("#DC2626")   # slightly deeper for light bg

PAGE_W, PAGE_H = A4
HEADER_H  = 2.8 * cm
FOOTER_H  = 0.9 * cm
MARGIN_LR = 1.5 * cm

MONTHS_ES       = ['Enero','Febrero','Marzo','Abril','Mayo','Junio',
                   'Julio','Agosto','Septiembre','Octubre','Noviembre','Diciembre']
MONTHS_ES_LOWER = [m.lower() for m in MONTHS_ES]

_TYPE_LABELS = {
    "ingresos":    "Reporte de Ingresos",
    "egresos":     "Reporte de Egresos",
    "ambos":       "Reporte Completo",
    "presupuesto": "Reporte de Presupuesto 50-30-20",
}


def _period_label(mes, ano):
    if mes and ano:
        return f"{MONTHS_ES[mes - 1]} {ano}"
    if ano:
        return str(ano)
    return "Histórico"


def _tipo_label(tipo):
    return "Ingreso" if tipo == "ingreso" else "Egreso"


def _fmt(amount):
    return f"S/. {amount:,.2f}"


def build_filename(report_type, mes, ano, ext):
    if mes and ano:
        return f"{MONTHS_ES_LOWER[mes - 1]}_{ano}_{report_type}_mensual.{ext}"
    if ano:
        return f"{ano}_{report_type}_anual.{ext}"
    return f"historico_{report_type}.{ext}"


def _draw_page(canv, doc, report_type, mes, ano):
    canv.saveState()

    header_y = PAGE_H - HEADER_H

    # ── Header background ──────────────────────────────────────────────────
    canv.setFillColor(BG_HEADER)
    canv.rect(0, header_y, PAGE_W, HEADER_H, fill=1, stroke=0)

    # Neon top stripe (3 px)
    canv.setFillColor(NEON)
    canv.rect(0, PAGE_H - 3, PAGE_W, 3, fill=1, stroke=0)

    # Bottom rule of header
    canv.setStrokeColor(RULE)
    canv.setLineWidth(0.5)
    canv.line(0, header_y, PAGE_W, header_y)

    # ── Brand: VAULT (ink) + LAS (light gray) ─────────────────────────────
    font_sz = 22
    vault_x = MARGIN_LR + 0.1 * cm
    vault_y = PAGE_H - 1.6 * cm

    canv.setFillColor(INK)
    canv.setFont("Helvetica-Bold", font_sz)
    canv.drawString(vault_x, vault_y, "VAULT")
    vault_w = canv.stringWidth("VAULT", "Helvetica-Bold", font_sz)

    canv.setFillColor(INK_LIGHT)
    canv.setFont("Helvetica", font_sz)
    canv.drawString(vault_x + vault_w, vault_y, "LAS")

    # Subtitle: report type
    canv.setFillColor(INK_LIGHT)
    canv.setFont("Helvetica", 7.5)
    canv.drawString(MARGIN_LR + 0.1 * cm, PAGE_H - 2.2 * cm,
                    _TYPE_LABELS.get(report_type, "Reporte").upper())

    # ── Period badge (neon pill on right) ─────────────────────────────────
    period  = _period_label(mes, ano)
    badge_w = 3.4 * cm
    badge_h = 0.65 * cm
    badge_x = PAGE_W - MARGIN_LR - badge_w
    badge_y = header_y + (HEADER_H - badge_h) / 2

    canv.setFillColor(NEON)
    canv.roundRect(badge_x, badge_y, badge_w, badge_h, 3, fill=1, stroke=0)
    canv.setFillColor(INK)
    canv.setFont("Helvetica-Bold", 8.5)
    canv.drawCentredString(badge_x + badge_w / 2, badge_y + 0.18 * cm, period)

    # Subtle decorative dots (right of subtitle, left of badge)
    for row in range(3):
        for col in range(6):
            canv.setFillColor(RULE)
            canv.circle(
                PAGE_W - MARGIN_LR - badge_w - 0.6 * cm - col * 0.32 * cm,
                header_y + 0.45 * cm + row * 0.32 * cm,
                0.04 * cm, fill=1, stroke=0,
            )

    # ── Footer ────────────────────────────────────────────────────────────
    canv.setFillColor(BG_FOOT)
    canv.rect(0, 0, PAGE_W, FOOTER_H, fill=1, stroke=0)

    # Neon top line of footer
    canv.setFillColor(NEON)
    canv.rect(0, FOOTER_H - 0.5, PAGE_W, 0.5, fill=1, stroke=0)

    canv.setFillColor(INK_LIGHT)
    canv.setFont("Helvetica", 6.5)
    canv.drawString(MARGIN_LR, 0.3 * cm,
                    f"Generado {datetime.now().strftime('%d/%m/%Y %H:%M')}")

    canv.setFillColor(INK_MID)
    canv.setFont("Helvetica-Bold", 7)
    canv.drawCentredString(PAGE_W / 2, 0.3 * cm, "VAULTLAS")

    canv.setFillColor(INK_LIGHT)
    canv.setFont("Helvetica", 6.5)
    canv.drawRightString(PAGE_W - MARGIN_LR, 0.3 * cm, f"Página {doc.page:02d}")

    canv.restoreState()


def generate_pdf_report(movements, report_type, mes=None, ano=None):
    buffer = BytesIO()

    margin_top = HEADER_H + 0.5 * cm
    margin_bot = FOOTER_H + 0.5 * cm

    doc = BaseDocTemplate(
        buffer, pagesize=A4,
        leftMargin=MARGIN_LR, rightMargin=MARGIN_LR,
        topMargin=margin_top, bottomMargin=margin_bot,
        title=f"Vaultlas — {_TYPE_LABELS.get(report_type, 'Reporte')} {_period_label(mes, ano)}",
        author="Vaultlas Finance",
        subject=_TYPE_LABELS.get(report_type, "Reporte"),
        creator="Vaultlas",
    )
    frame = Frame(
        MARGIN_LR, margin_bot,
        PAGE_W - 2 * MARGIN_LR,
        PAGE_H - margin_top - margin_bot,
        id="body", showBoundary=0,
    )
    doc.addPageTemplates([
        PageTemplate(id="main", frames=[frame],
                     onPage=lambda c, d: _draw_page(c, d, report_type, mes, ano))
    ])

    styles = getSampleStyleSheet()

    # ── Stats bar ─────────────────────────────────────────────────────────
    total    = sum(float(m["monto"]) for m in movements)
    ingresos = sum(float(m["monto"]) for m in movements if m["tipo"] == "ingreso")
    egresos  = sum(float(m["monto"]) for m in movements if m["tipo"] == "egreso")

    if report_type == "ingresos":
        stats = [
            ("MOVIMIENTOS",    str(len(movements)), INK),
            ("TOTAL INGRESOS", _fmt(total),         EMERALD),
        ]
    elif report_type == "egresos":
        stats = [
            ("MOVIMIENTOS",   str(len(movements)), INK),
            ("TOTAL EGRESOS", _fmt(total),         CORAL),
        ]
    else:
        neto = ingresos - egresos
        stats = [
            ("MOVIMIENTOS", str(len(movements)), INK),
            ("INGRESOS",    _fmt(ingresos),      EMERALD),
            ("EGRESOS",     _fmt(egresos),        CORAL),
            ("NETO",        _fmt(neto),           EMERALD if neto >= 0 else CORAL),
        ]

    # ── Stats bar — single row, label stacked above value ────────────────
    stat_col_w = (PAGE_W - 2 * MARGIN_LR) / len(stats)

    def _hex(color_obj):
        """Convert reportlab color to hex string for Paragraph XML."""
        return f"#{color_obj.hexval()}"

    stat_cells = []
    for label, value, val_color in stats:
        cell_style = ParagraphStyle(
            f"Stat_{label}",
            parent=styles["Normal"],
            alignment=TA_CENTER,
            leading=20,
            spaceAfter=0,
        )
        markup = (
            f'<font name="Helvetica" size="7" color="{_hex(INK_LIGHT)}">'
            f'{label}</font><br/>'
            f'<font name="Helvetica-Bold" size="12" color="{_hex(val_color)}">'
            f'{value}</font>'
        )
        stat_cells.append(Paragraph(markup, cell_style))

    stat_table = Table(
        [stat_cells],
        colWidths=[stat_col_w] * len(stats),
        rowHeights=[1.5 * cm],
    )
    stat_table.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), BG_STATS),
        ("TOPPADDING",    (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING",   (0, 0), (-1, -1), 8),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 8),
        ("LINEAFTER",     (0, 0), (-2, -1), 0.5, RULE),
        ("BOX",           (0, 0), (-1, -1), 0.5, RULE),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
    ]))

    # ── Movements table ───────────────────────────────────────────────────
    if report_type == "ambos":
        headers    = ["Fecha", "Descripción", "N° Op.", "Tipo", "Monto"]
        col_widths = [2.0 * cm, 7.8 * cm, 2.5 * cm, 2.0 * cm, 2.8 * cm]
    else:
        headers    = ["Fecha", "Descripción", "N° Op.", "Monto"]
        col_widths = [2.0 * cm, 10.3 * cm, 2.5 * cm, 2.8 * cm]

    data = [headers]
    for m in movements:
        monto  = float(m["monto"])
        num_op = m.get("codigo_operacion") or "—"
        if report_type == "ambos":
            data.append([m["fecha"], m["descripcion"], num_op,
                         _tipo_label(m["tipo"]), _fmt(monto)])
        else:
            data.append([m["fecha"], m["descripcion"], num_op, _fmt(monto)])

    total_label = {"ingresos": "Total Ingresos", "egresos": "Total Egresos"}.get(report_type, "Total")
    if report_type == "ambos":
        data.append(["", "", "", total_label, _fmt(total)])
    else:
        data.append(["", "", total_label, _fmt(total)])

    n = len(data)
    mv_table = Table(data, colWidths=col_widths, repeatRows=1)

    ts = TableStyle([
        # Header row
        ("BACKGROUND",     (0, 0),  (-1, 0),      INK),
        ("TEXTCOLOR",      (0, 0),  (-1, 0),      NEON),
        ("FONTNAME",       (0, 0),  (-1, 0),      "Helvetica-Bold"),
        ("FONTSIZE",       (0, 0),  (-1, 0),      8),
        ("ALIGN",          (0, 0),  (-1, 0),      "CENTER"),
        ("LINEBELOW",      (0, 0),  (-1, 0),      1.5, NEON),
        # Data rows
        ("ROWBACKGROUNDS", (0, 1),  (-1, n - 2),  [BG_ROW_MAIN, BG_ROW_ALT]),
        ("TEXTCOLOR",      (0, 1),  (-1, n - 2),  INK),
        ("FONTNAME",       (0, 1),  (-1, n - 2),  "Helvetica"),
        ("FONTSIZE",       (0, 1),  (-1, n - 2),  8),
        ("LINEBELOW",      (0, 1),  (-1, n - 2),  0.25, RULE),
        # Total row
        ("BACKGROUND",     (0, n - 1), (-1, n - 1), BG_STATS),
        ("TEXTCOLOR",      (0, n - 1), (-1, n - 1), INK),
        ("FONTNAME",       (0, n - 1), (-1, n - 1), "Helvetica-Bold"),
        ("FONTSIZE",       (0, n - 1), (-1, n - 1), 8.5),
        ("LINEABOVE",      (0, n - 1), (-1, n - 1), 1, RULE),
        # Alignment
        ("ALIGN",          (-1, 0), (-1, -1), "RIGHT"),
        ("ALIGN",          (0, 1),  (0, -1),  "CENTER"),
        ("VALIGN",         (0, 0),  (-1, -1), "MIDDLE"),
        # Padding
        ("TOPPADDING",     (0, 0),  (-1, -1), 5),
        ("BOTTOMPADDING",  (0, 0),  (-1, -1), 5),
        ("LEFTPADDING",    (0, 0),  (-1, -1), 6),
        ("RIGHTPADDING",   (0, 0),  (-1, -1), 6),
        # Outer box
        ("BOX",            (0, 0),  (-1, -1), 0.5, RULE),
    ])

    # Color amounts and tipo column
    if report_type == "ambos":
        for i, m in enumerate(movements, start=1):
            c = EMERALD if m["tipo"] == "ingreso" else CORAL
            ts.add("TEXTCOLOR", (-1, i), (-1, i), c)
            ts.add("FONTNAME",  (-1, i), (-1, i), "Helvetica-Bold")
            ts.add("TEXTCOLOR", (-2, i), (-2, i), c)
    else:
        c = EMERALD if report_type == "ingresos" else CORAL
        ts.add("TEXTCOLOR", (-1, 1), (-1, n - 2), c)
        ts.add("FONTNAME",  (-1, 1), (-1, n - 2), "Helvetica-Bold")

    mv_table.setStyle(ts)
    doc.build([stat_table, Spacer(1, 0.45 * cm), mv_table])
    buffer.seek(0)
    return buffer.read()


def generate_excel_report(movements, report_type, mes=None, ano=None):
    wb = openpyxl.Workbook()
    ws = wb.active

    type_labels = {"ingresos": "Ingresos", "egresos": "Egresos", "ambos": "Completo"}
    ws.title = type_labels.get(report_type, "Reporte")

    # Light palette
    NEON_HEX     = "CCFF00"
    INK_HEX      = "111111"
    INK_MID_HEX  = "555555"
    INK_LITE_HEX = "999999"
    BG_HEAD_HEX  = "F9F9F9"
    BG_STAT_HEX  = "F4F4F4"
    BG_ROW_HEX   = "FFFFFF"
    BG_ALT_HEX   = "FAFAFA"
    RULE_HEX     = "E5E5E5"
    EMERALD_HEX  = "059669"
    CORAL_HEX    = "DC2626"

    period_label    = _period_label(mes, ano)
    type_label_full = _TYPE_LABELS.get(report_type, "Reporte")

    n_cols   = 5 if report_type == "ambos" else 4
    last_col = get_column_letter(n_cols)

    head_fill = PatternFill(start_color=BG_HEAD_HEX, end_color=BG_HEAD_HEX, fill_type="solid")

    # Row 1: brand
    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value     = "VAULTLAS"
    c.font      = Font(name="Calibri", bold=True, color=INK_HEX, size=18)
    c.fill      = head_fill
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32

    # Row 2: report type
    ws.merge_cells(f"A2:{last_col}2")
    c = ws["A2"]
    c.value     = type_label_full
    c.font      = Font(name="Calibri", color=INK_MID_HEX, size=11)
    c.fill      = head_fill
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 18

    # Row 3: period + generated
    ws.merge_cells(f"A3:{last_col}3")
    c = ws["A3"]
    c.value     = f"Período: {period_label}    ·    Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    c.font      = Font(name="Calibri", color=INK_LITE_HEX, size=9)
    c.fill      = head_fill
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 16

    # Row 4: neon stripe
    ws.merge_cells(f"A4:{last_col}4")
    ws["A4"].fill = PatternFill(start_color=NEON_HEX, end_color=NEON_HEX, fill_type="solid")
    ws.row_dimensions[4].height = 3

    # ── Stats (rows 5-6) ───────────────────────────────────────────────────
    total    = sum(float(m["monto"]) for m in movements)
    ingresos = sum(float(m["monto"]) for m in movements if m["tipo"] == "ingreso")
    egresos  = sum(float(m["monto"]) for m in movements if m["tipo"] == "egreso")

    if report_type == "ingresos":
        stat_pairs = [("MOVIMIENTOS", str(len(movements)), INK_HEX),
                      ("TOTAL INGRESOS", _fmt(total), EMERALD_HEX)]
    elif report_type == "egresos":
        stat_pairs = [("MOVIMIENTOS", str(len(movements)), INK_HEX),
                      ("TOTAL EGRESOS", _fmt(total), CORAL_HEX)]
    else:
        neto = ingresos - egresos
        stat_pairs = [
            ("MOVIMIENTOS", str(len(movements)), INK_HEX),
            ("INGRESOS",    _fmt(ingresos),      EMERALD_HEX),
            ("EGRESOS",     _fmt(egresos),        CORAL_HEX),
            ("NETO",        _fmt(neto),           EMERALD_HEX if neto >= 0 else CORAL_HEX),
        ]

    stat_fill = PatternFill(start_color=BG_STAT_HEX, end_color=BG_STAT_HEX, fill_type="solid")
    rule_border = Border(
        bottom=Side(style="thin", color=RULE_HEX),
        right=Side(style="thin",  color=RULE_HEX),
    )
    ws.row_dimensions[5].height = 14
    ws.row_dimensions[6].height = 22

    for col_idx, (lbl, val, val_color) in enumerate(stat_pairs, start=1):
        lc = ws.cell(row=5, column=col_idx, value=lbl)
        lc.font      = Font(name="Calibri", color=INK_LITE_HEX, size=8)
        lc.fill      = stat_fill
        lc.alignment = Alignment(horizontal="center", vertical="bottom")
        lc.border    = rule_border

        vc = ws.cell(row=6, column=col_idx, value=val)
        vc.font      = Font(name="Calibri", bold=True, color=val_color, size=11)
        vc.fill      = stat_fill
        vc.alignment = Alignment(horizontal="center", vertical="center")
        vc.border    = rule_border

    # Row 7: blank separator
    for col_idx in range(1, n_cols + 1):
        ws.cell(row=7, column=col_idx).fill = \
            PatternFill(start_color=RULE_HEX, end_color=RULE_HEX, fill_type="solid")
    ws.row_dimensions[7].height = 4

    # ── Column headers (row 8) ─────────────────────────────────────────────
    header_row = 8
    if report_type == "ambos":
        headers    = ["Fecha", "Descripción", "N° Operación", "Tipo", "Monto"]
        col_widths = [14, 45, 18, 12, 16]
    else:
        headers    = ["Fecha", "Descripción", "N° Operación", "Monto"]
        col_widths = [14, 55, 18, 16]

    hdr_fill = PatternFill(start_color=INK_HEX, end_color=INK_HEX, fill_type="solid")
    hdr_font = Font(name="Calibri", bold=True, color=NEON_HEX, size=10)

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill      = hdr_fill
        cell.font      = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = rule_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[header_row].height = 22

    # ── Data rows ─────────────────────────────────────────────────────────
    row_fill_a   = PatternFill(start_color=BG_ROW_HEX, end_color=BG_ROW_HEX, fill_type="solid")
    row_fill_b   = PatternFill(start_color=BG_ALT_HEX, end_color=BG_ALT_HEX, fill_type="solid")
    default_font = Font(name="Calibri", color=INK_HEX, size=10)

    running_total = 0.0
    for i, m in enumerate(movements):
        data_row = header_row + 1 + i
        monto  = float(m["monto"])
        running_total += monto
        num_op = m.get("codigo_operacion") or "—"
        fill   = row_fill_a if i % 2 == 0 else row_fill_b

        if report_type == "ambos":
            row_data = [m["fecha"], m["descripcion"], num_op,
                        _tipo_label(m["tipo"]), monto]
        else:
            row_data = [m["fecha"], m["descripcion"], num_op, monto]

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=data_row, column=col_idx, value=value)
            cell.fill   = fill
            cell.border = rule_border
            cell.alignment = Alignment(vertical="center")

            is_amount = col_idx == len(row_data)
            is_tipo   = report_type == "ambos" and col_idx == 4

            if is_amount:
                clr = (EMERALD_HEX if m["tipo"] == "ingreso" else CORAL_HEX) \
                      if report_type == "ambos" \
                      else (EMERALD_HEX if report_type == "ingresos" else CORAL_HEX)
                cell.font          = Font(name="Calibri", bold=True, color=clr, size=10)
                cell.number_format = '"S/. "#,##0.00'
                cell.alignment     = Alignment(horizontal="right", vertical="center")
            elif is_tipo:
                clr       = EMERALD_HEX if m["tipo"] == "ingreso" else CORAL_HEX
                cell.font = Font(name="Calibri", color=clr, size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = default_font

        ws.row_dimensions[data_row].height = 16

    # ── Total row ─────────────────────────────────────────────────────────
    total_row   = header_row + 1 + len(movements)
    total_label = {"ingresos": "Total Ingresos", "egresos": "Total Egresos"}.get(report_type, "Total")
    total_fill  = PatternFill(start_color=BG_STAT_HEX, end_color=BG_STAT_HEX, fill_type="solid")
    total_font  = Font(name="Calibri", bold=True, color=INK_HEX, size=11)

    lbl_col = len(headers) - 1
    val_col = len(headers)

    lbl_c = ws.cell(row=total_row, column=lbl_col, value=total_label)
    lbl_c.fill      = total_fill
    lbl_c.font      = total_font
    lbl_c.alignment = Alignment(horizontal="right", vertical="center")
    lbl_c.border    = rule_border

    val_c = ws.cell(row=total_row, column=val_col, value=running_total)
    val_c.fill          = total_fill
    val_c.font          = Font(name="Calibri", bold=True, color=INK_HEX, size=11)
    val_c.number_format = '"S/. "#,##0.00'
    val_c.alignment     = Alignment(horizontal="right", vertical="center")
    val_c.border        = rule_border

    for col_idx in range(1, lbl_col):
        c = ws.cell(row=total_row, column=col_idx)
        c.fill   = total_fill
        c.border = rule_border

    ws.row_dimensions[total_row].height = 22

    # Freeze panes + hide gridlines
    ws.freeze_panes    = ws.cell(row=header_row + 1, column=1)
    ws.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


# ── Budget 50-30-20 exports ───────────────────────────────────────────────────

def generate_budget_pdf_report(budgets_data: list, ano: int, mes: int = None):
    """
    budgets_data: list of dicts with keys:
      month_label (str, e.g. "Enero"), total_income, budgeted_needs, actual_needs,
      budgeted_wants, actual_wants, budgeted_savings, actual_savings,
      actual_debt, net (float)
    """
    buffer = BytesIO()

    margin_top = HEADER_H + 0.5 * cm
    margin_bot = FOOTER_H + 0.5 * cm
    report_type = "presupuesto"

    doc = BaseDocTemplate(
        buffer, pagesize=A4,
        leftMargin=MARGIN_LR, rightMargin=MARGIN_LR,
        topMargin=margin_top, bottomMargin=margin_bot,
        title=f"Vaultlas — {_TYPE_LABELS[report_type]} {_period_label(mes, ano)}",
        author="Vaultlas Finance",
        subject=_TYPE_LABELS[report_type],
        creator="Vaultlas",
    )
    frame = Frame(
        MARGIN_LR, margin_bot,
        PAGE_W - 2 * MARGIN_LR,
        PAGE_H - margin_top - margin_bot,
        id="body", showBoundary=0,
    )
    doc.addPageTemplates([
        PageTemplate(id="main", frames=[frame],
                     onPage=lambda c, d: _draw_page(c, d, report_type, mes, ano))
    ])

    styles = getSampleStyleSheet()

    # ── Stats bar ─────────────────────────────────────────────────────────
    total_income   = sum(r["total_income"] for r in budgets_data)
    total_expenses = sum(
        r["actual_needs"] + r["actual_wants"] + r["actual_savings"] + r["actual_debt"]
        for r in budgets_data
    )
    total_net = total_income - total_expenses

    period_label = str(ano) if not mes else f"{MONTHS_ES[mes - 1]} {ano}"
    stats = [
        ("PERÍODO",         period_label,       INK),
        ("INGRESOS TOTAL",  _fmt(total_income),  EMERALD),
        ("GASTOS TOTAL",    _fmt(total_expenses), CORAL),
        ("NETO",            _fmt(total_net),      EMERALD if total_net >= 0 else CORAL),
    ]

    stat_col_w = (PAGE_W - 2 * MARGIN_LR) / len(stats)

    def _hex(color_obj):
        return f"#{color_obj.hexval()}"

    stat_cells = []
    for label, value, val_color in stats:
        cell_style = ParagraphStyle(
            f"Stat_{label}",
            parent=styles["Normal"],
            alignment=TA_CENTER,
            leading=20,
            spaceAfter=0,
        )
        markup = (
            f'<font name="Helvetica" size="7" color="{_hex(INK_LIGHT)}">'
            f'{label}</font><br/>'
            f'<font name="Helvetica-Bold" size="12" color="{_hex(val_color)}">'
            f'{value}</font>'
        )
        stat_cells.append(Paragraph(markup, cell_style))

    stat_table = Table(
        [stat_cells],
        colWidths=[stat_col_w] * len(stats),
        rowHeights=[1.5 * cm],
    )
    stat_table.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), BG_STATS),
        ("TOPPADDING",    (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING",   (0, 0), (-1, -1), 8),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 8),
        ("LINEAFTER",     (0, 0), (-2, -1), 0.5, RULE),
        ("BOX",           (0, 0), (-1, -1), 0.5, RULE),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
    ]))

    # ── Budget table ──────────────────────────────────────────────────────
    col_widths = [2.2 * cm, 2.8 * cm, 2.2 * cm, 1.4 * cm,
                  2.2 * cm, 1.4 * cm, 2.4 * cm, 2.0 * cm]
    headers = ["Mes", "Ingresos", "Nec. (S/.)", "Nec. %",
               "Gustos (S/.)", "Gustos %", "Ahorros (S/.)", "Neto"]

    def _pct(actual, budgeted):
        if budgeted and budgeted > 0:
            return actual / budgeted * 100
        return 0.0

    data = [headers]
    pct_cells_green = []   # (row_idx, col_idx) that should be green
    pct_cells_red   = []   # (row_idx, col_idx) that should be red

    for row_idx, r in enumerate(budgets_data, start=1):
        needs_pct   = _pct(r["actual_needs"],   r["budgeted_needs"])
        wants_pct   = _pct(r["actual_wants"],   r["budgeted_wants"])

        needs_pct_str  = f"{needs_pct:.0f}%"
        wants_pct_str  = f"{wants_pct:.0f}%"

        row_data = [
            r["month_label"],
            _fmt(r["total_income"]),
            _fmt(r["actual_needs"]),
            needs_pct_str,
            _fmt(r["actual_wants"]),
            wants_pct_str,
            _fmt(r["actual_savings"]),
            _fmt(r["net"]),
        ]
        data.append(row_data)

        # Track percentage column colors (col indices 3 and 5, 0-based)
        if needs_pct <= 100:
            pct_cells_green.append((row_idx, 3))
        else:
            pct_cells_red.append((row_idx, 3))

        if wants_pct <= 100:
            pct_cells_green.append((row_idx, 5))
        else:
            pct_cells_red.append((row_idx, 5))

    # Total row
    tot_income   = sum(r["total_income"]   for r in budgets_data)
    tot_needs    = sum(r["actual_needs"]   for r in budgets_data)
    tot_wants    = sum(r["actual_wants"]   for r in budgets_data)
    tot_savings  = sum(r["actual_savings"] for r in budgets_data)
    tot_net      = sum(r["net"]            for r in budgets_data)
    bud_needs    = sum(r["budgeted_needs"] for r in budgets_data)
    bud_wants    = sum(r["budgeted_wants"] for r in budgets_data)
    tot_needs_pct  = f"{_pct(tot_needs,  bud_needs):.0f}%"
    tot_wants_pct  = f"{_pct(tot_wants,  bud_wants):.0f}%"

    data.append([
        "TOTAL",
        _fmt(tot_income),
        _fmt(tot_needs),
        tot_needs_pct,
        _fmt(tot_wants),
        tot_wants_pct,
        _fmt(tot_savings),
        _fmt(tot_net),
    ])

    n = len(data)
    bud_table = Table(data, colWidths=col_widths, repeatRows=1)

    ts = TableStyle([
        # Header row
        ("BACKGROUND",     (0, 0),  (-1, 0),      INK),
        ("TEXTCOLOR",      (0, 0),  (-1, 0),      NEON),
        ("FONTNAME",       (0, 0),  (-1, 0),      "Helvetica-Bold"),
        ("FONTSIZE",       (0, 0),  (-1, 0),      8),
        ("ALIGN",          (0, 0),  (-1, 0),      "CENTER"),
        ("LINEBELOW",      (0, 0),  (-1, 0),      1.5, NEON),
        # Data rows
        ("ROWBACKGROUNDS", (0, 1),  (-1, n - 2),  [BG_ROW_MAIN, BG_ROW_ALT]),
        ("TEXTCOLOR",      (0, 1),  (-1, n - 2),  INK),
        ("FONTNAME",       (0, 1),  (-1, n - 2),  "Helvetica"),
        ("FONTSIZE",       (0, 1),  (-1, n - 2),  8),
        ("LINEBELOW",      (0, 1),  (-1, n - 2),  0.25, RULE),
        # Total row
        ("BACKGROUND",     (0, n - 1), (-1, n - 1), BG_STATS),
        ("TEXTCOLOR",      (0, n - 1), (-1, n - 1), INK),
        ("FONTNAME",       (0, n - 1), (-1, n - 1), "Helvetica-Bold"),
        ("FONTSIZE",       (0, n - 1), (-1, n - 1), 8.5),
        ("LINEABOVE",      (0, n - 1), (-1, n - 1), 1, RULE),
        # Alignment: right-align numeric columns
        ("ALIGN",          (1, 1),  (1, -1),  "RIGHT"),
        ("ALIGN",          (2, 1),  (2, -1),  "RIGHT"),
        ("ALIGN",          (3, 0),  (3, -1),  "CENTER"),
        ("ALIGN",          (4, 1),  (4, -1),  "RIGHT"),
        ("ALIGN",          (5, 0),  (5, -1),  "CENTER"),
        ("ALIGN",          (6, 1),  (6, -1),  "RIGHT"),
        ("ALIGN",          (7, 1),  (7, -1),  "RIGHT"),
        ("ALIGN",          (0, 0),  (0, -1),  "LEFT"),
        ("VALIGN",         (0, 0),  (-1, -1), "MIDDLE"),
        # Padding
        ("TOPPADDING",     (0, 0),  (-1, -1), 5),
        ("BOTTOMPADDING",  (0, 0),  (-1, -1), 5),
        ("LEFTPADDING",    (0, 0),  (-1, -1), 4),
        ("RIGHTPADDING",   (0, 0),  (-1, -1), 4),
        # Outer box
        ("BOX",            (0, 0),  (-1, -1), 0.5, RULE),
    ])

    # Apply percentage column colors
    for (ri, ci) in pct_cells_green:
        ts.add("TEXTCOLOR", (ci, ri), (ci, ri), EMERALD)
        ts.add("FONTNAME",  (ci, ri), (ci, ri), "Helvetica-Bold")
    for (ri, ci) in pct_cells_red:
        ts.add("TEXTCOLOR", (ci, ri), (ci, ri), CORAL)
        ts.add("FONTNAME",  (ci, ri), (ci, ri), "Helvetica-Bold")

    # Color Neto column: green if positive, red if negative
    for row_idx, r in enumerate(budgets_data, start=1):
        c = EMERALD if r["net"] >= 0 else CORAL
        ts.add("TEXTCOLOR", (7, row_idx), (7, row_idx), c)
        ts.add("FONTNAME",  (7, row_idx), (7, row_idx), "Helvetica-Bold")

    bud_table.setStyle(ts)
    doc.build([stat_table, Spacer(1, 0.45 * cm), bud_table])
    buffer.seek(0)
    return buffer.read()


def generate_budget_excel_report(budgets_data: list, ano: int, mes: int = None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Presupuesto"

    # Palette hex strings (same as generate_excel_report)
    NEON_HEX     = "CCFF00"
    INK_HEX      = "111111"
    INK_MID_HEX  = "555555"
    INK_LITE_HEX = "999999"
    BG_HEAD_HEX  = "F9F9F9"
    BG_STAT_HEX  = "F4F4F4"
    BG_ROW_HEX   = "FFFFFF"
    BG_ALT_HEX   = "FAFAFA"
    RULE_HEX     = "E5E5E5"
    EMERALD_HEX  = "059669"
    CORAL_HEX    = "DC2626"

    period_label    = _period_label(mes, ano)
    type_label_full = _TYPE_LABELS["presupuesto"]

    n_cols   = 8
    last_col = get_column_letter(n_cols)

    head_fill = PatternFill(start_color=BG_HEAD_HEX, end_color=BG_HEAD_HEX, fill_type="solid")

    # Row 1: brand
    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value     = "VAULTLAS"
    c.font      = Font(name="Calibri", bold=True, color=INK_HEX, size=18)
    c.fill      = head_fill
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32

    # Row 2: report type
    ws.merge_cells(f"A2:{last_col}2")
    c = ws["A2"]
    c.value     = type_label_full
    c.font      = Font(name="Calibri", color=INK_MID_HEX, size=11)
    c.fill      = head_fill
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 18

    # Row 3: period + generated
    ws.merge_cells(f"A3:{last_col}3")
    c = ws["A3"]
    c.value     = f"Período: {period_label}    ·    Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    c.font      = Font(name="Calibri", color=INK_LITE_HEX, size=9)
    c.fill      = head_fill
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 16

    # Row 4: neon stripe
    ws.merge_cells(f"A4:{last_col}4")
    ws["A4"].fill = PatternFill(start_color=NEON_HEX, end_color=NEON_HEX, fill_type="solid")
    ws.row_dimensions[4].height = 3

    # ── Stats (rows 5-6) ───────────────────────────────────────────────────
    total_income   = sum(r["total_income"] for r in budgets_data)
    total_expenses = sum(
        r["actual_needs"] + r["actual_wants"] + r["actual_savings"] + r["actual_debt"]
        for r in budgets_data
    )
    total_net = total_income - total_expenses

    stat_pairs = [
        ("PERÍODO",        period_label,          INK_HEX),
        ("INGRESOS TOTAL", _fmt(total_income),     EMERALD_HEX),
        ("GASTOS TOTAL",   _fmt(total_expenses),   CORAL_HEX),
        ("NETO",           _fmt(total_net),         EMERALD_HEX if total_net >= 0 else CORAL_HEX),
    ]

    stat_fill = PatternFill(start_color=BG_STAT_HEX, end_color=BG_STAT_HEX, fill_type="solid")
    rule_border = Border(
        bottom=Side(style="thin", color=RULE_HEX),
        right=Side(style="thin",  color=RULE_HEX),
    )
    ws.row_dimensions[5].height = 14
    ws.row_dimensions[6].height = 22

    for col_idx, (lbl, val, val_color) in enumerate(stat_pairs, start=1):
        lc = ws.cell(row=5, column=col_idx, value=lbl)
        lc.font      = Font(name="Calibri", color=INK_LITE_HEX, size=8)
        lc.fill      = stat_fill
        lc.alignment = Alignment(horizontal="center", vertical="bottom")
        lc.border    = rule_border

        vc = ws.cell(row=6, column=col_idx, value=val)
        vc.font      = Font(name="Calibri", bold=True, color=val_color, size=11)
        vc.fill      = stat_fill
        vc.alignment = Alignment(horizontal="center", vertical="center")
        vc.border    = rule_border

    # Fill remaining stat cols (5-8)
    for col_idx in range(5, n_cols + 1):
        for r in (5, 6):
            c2 = ws.cell(row=r, column=col_idx)
            c2.fill   = stat_fill
            c2.border = rule_border

    # Row 7: blank separator
    for col_idx in range(1, n_cols + 1):
        ws.cell(row=7, column=col_idx).fill = \
            PatternFill(start_color=RULE_HEX, end_color=RULE_HEX, fill_type="solid")
    ws.row_dimensions[7].height = 4

    # ── Column headers (row 8) ─────────────────────────────────────────────
    header_row = 8
    headers    = ["Mes", "Ingresos", "Nec. (S/.)", "Nec. %",
                  "Gustos (S/.)", "Gustos %", "Ahorros (S/.)", "Neto"]
    col_widths = [14, 16, 14, 10, 14, 10, 14, 14]

    hdr_fill = PatternFill(start_color=INK_HEX, end_color=INK_HEX, fill_type="solid")
    hdr_font = Font(name="Calibri", bold=True, color=NEON_HEX, size=10)

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill      = hdr_fill
        cell.font      = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = rule_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[header_row].height = 22

    # ── Data rows ─────────────────────────────────────────────────────────
    row_fill_a   = PatternFill(start_color=BG_ROW_HEX, end_color=BG_ROW_HEX, fill_type="solid")
    row_fill_b   = PatternFill(start_color=BG_ALT_HEX, end_color=BG_ALT_HEX, fill_type="solid")
    default_font = Font(name="Calibri", color=INK_HEX, size=10)
    money_fmt    = '"S/. "#,##0.00'

    def _pct(actual, budgeted):
        if budgeted and budgeted > 0:
            return actual / budgeted * 100
        return 0.0

    for i, r in enumerate(budgets_data):
        data_row = header_row + 1 + i
        fill     = row_fill_a if i % 2 == 0 else row_fill_b

        needs_pct  = _pct(r["actual_needs"],  r["budgeted_needs"])
        wants_pct  = _pct(r["actual_wants"],  r["budgeted_wants"])

        row_values = [
            r["month_label"],
            r["total_income"],
            r["actual_needs"],
            f"{needs_pct:.0f}%",
            r["actual_wants"],
            f"{wants_pct:.0f}%",
            r["actual_savings"],
            r["net"],
        ]

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=data_row, column=col_idx, value=value)
            cell.fill   = fill
            cell.border = rule_border

            # Percentage columns (4 = Nec.%, 6 = Gustos%)
            if col_idx == 4:
                clr = EMERALD_HEX if needs_pct <= 100 else CORAL_HEX
                cell.font      = Font(name="Calibri", bold=True, color=clr, size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                clr = EMERALD_HEX if wants_pct <= 100 else CORAL_HEX
                cell.font      = Font(name="Calibri", bold=True, color=clr, size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (2, 3, 5, 7):
                # Money columns
                cell.font          = Font(name="Calibri", color=INK_HEX, size=10)
                cell.number_format = money_fmt
                cell.alignment     = Alignment(horizontal="right", vertical="center")
            elif col_idx == 8:
                # Neto
                clr = EMERALD_HEX if r["net"] >= 0 else CORAL_HEX
                cell.font          = Font(name="Calibri", bold=True, color=clr, size=10)
                cell.number_format = money_fmt
                cell.alignment     = Alignment(horizontal="right", vertical="center")
            else:
                cell.font      = default_font
                cell.alignment = Alignment(vertical="center")

        ws.row_dimensions[data_row].height = 16

    # ── Total row ─────────────────────────────────────────────────────────
    total_row  = header_row + 1 + len(budgets_data)
    total_fill = PatternFill(start_color=BG_STAT_HEX, end_color=BG_STAT_HEX, fill_type="solid")
    total_font = Font(name="Calibri", bold=True, color=INK_HEX, size=11)

    tot_income   = sum(r["total_income"]   for r in budgets_data)
    tot_needs    = sum(r["actual_needs"]   for r in budgets_data)
    tot_wants    = sum(r["actual_wants"]   for r in budgets_data)
    tot_savings  = sum(r["actual_savings"] for r in budgets_data)
    tot_net      = sum(r["net"]            for r in budgets_data)
    bud_needs    = sum(r["budgeted_needs"] for r in budgets_data)
    bud_wants    = sum(r["budgeted_wants"] for r in budgets_data)
    tot_needs_pct = _pct(tot_needs, bud_needs)
    tot_wants_pct = _pct(tot_wants, bud_wants)

    total_values = [
        "TOTAL",
        tot_income,
        tot_needs,
        f"{tot_needs_pct:.0f}%",
        tot_wants,
        f"{tot_wants_pct:.0f}%",
        tot_savings,
        tot_net,
    ]

    for col_idx, value in enumerate(total_values, start=1):
        cell = ws.cell(row=total_row, column=col_idx, value=value)
        cell.fill   = total_fill
        cell.border = rule_border

        if col_idx == 4:
            clr = EMERALD_HEX if tot_needs_pct <= 100 else CORAL_HEX
            cell.font      = Font(name="Calibri", bold=True, color=clr, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx == 6:
            clr = EMERALD_HEX if tot_wants_pct <= 100 else CORAL_HEX
            cell.font      = Font(name="Calibri", bold=True, color=clr, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in (2, 3, 5, 7):
            cell.font          = total_font
            cell.number_format = money_fmt
            cell.alignment     = Alignment(horizontal="right", vertical="center")
        elif col_idx == 8:
            clr = EMERALD_HEX if tot_net >= 0 else CORAL_HEX
            cell.font          = Font(name="Calibri", bold=True, color=clr, size=11)
            cell.number_format = money_fmt
            cell.alignment     = Alignment(horizontal="right", vertical="center")
        else:
            cell.font      = total_font
            cell.alignment = Alignment(vertical="center")

    ws.row_dimensions[total_row].height = 22

    # Freeze panes + hide gridlines
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
